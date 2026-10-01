from datetime import timedelta
import hashlib
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd

from listensignal import build_pulse, cluster_topics, demo_workspace, detect_spikes, pulse_html, pulse_xlsx, rising_terms
from listensignal.analysis import WeekWindow, daily_counts, mention_table
from listensignal.demo import DEMO_BRANDS, DEMO_END, DEMO_START, SPIKE_DAYS, make_demo_articles
from listensignal.pulse import safe_cell


def _fingerprint(frame: pd.DataFrame) -> str:
    text = frame[["title", "summary", "source", "url", "sentiment"]].to_csv(index=False)
    return hashlib.sha256((text + frame["published"].astype(str).str.cat()).encode()).hexdigest()


def test_demo_is_deterministic_and_fictional():
    first, second = make_demo_articles(), make_demo_articles()
    assert _fingerprint(first) == _fingerprint(second)
    assert 550 <= len(first) <= 650
    assert first["url"].str.startswith("https://demo.listensignal.invalid/").all()
    assert first["title"].is_unique  # so the demo survives headline de-duplication
    published = first["published"].dt.date
    assert published.min() == DEMO_START and published.max() == DEMO_END
    assert first["sentiment_scorer"].eq("lexicon-v1").all()
    assert _fingerprint(make_demo_articles(seed=7)) != _fingerprint(first)


def test_demo_contains_the_engineered_spike_and_excludes_decoys():
    articles = make_demo_articles()
    mentions = mention_table(articles, DEMO_BRANDS)
    spikes = detect_spikes(daily_counts(mentions, [b.name for b in DEMO_BRANDS]))
    fjellbrus_spikes = set(spikes.loc[spikes["brand"] == "Fjellbrus", "date"])
    assert min(SPIKE_DAYS) in fjellbrus_spikes
    aurora = articles.loc[articles["title"].str.contains("nordlys", case=False) & ~articles["title"].str.contains("Energi")]
    assert len(aurora) > 0
    assert not mentions["article_id"].isin(aurora["article_id"]).any()


def test_topics_are_deterministic_and_exclude_brand_names():
    mentions = mention_table(make_demo_articles(), DEMO_BRANDS).drop_duplicates("article_id")
    texts = (mentions["title"] + ". " + mentions["summary"]).tolist()
    aliases = tuple(a for b in DEMO_BRANDS for a in b.aliases)
    one = cluster_topics(texts, k=8, remove_terms=aliases)
    two = cluster_topics(texts, k=8, remove_terms=aliases)
    assert (one.assignments == two.assignments).all()
    assert len(one.clusters) == 8 and one.clusters["size"].sum() == len(texts)
    terms = " ".join(one.clusters["top_terms"]).lower()
    assert "fjellbrus" not in terms and "kystkraft" not in terms


def test_rising_terms_find_the_recall_story():
    mentions = mention_table(make_demo_articles(), DEMO_BRANDS)
    window = WeekWindow.ending(DEMO_END)
    this = mentions.loc[mentions["date"] >= window.this_start]
    prev = mentions.loc[(mentions["date"] >= window.prev_start) & (mentions["date"] <= window.prev_end)]
    rising = rising_terms((this["title"] + ". " + this["summary"]).tolist(), (prev["title"] + ". " + prev["summary"]).tolist())
    assert rising["term"].str.startswith("tilbakekall").any()


def test_pulse_exports_xlsx_and_html_with_limits_and_demo_notice():
    report = build_pulse(demo_workspace())
    assert report.window.this_end == DEMO_END
    assert any("Fjellbrus" in note and "spike" in note for note in report.headlines)
    workbook = load_workbook(BytesIO(pulse_xlsx(report)))
    assert {"About", "What changed", "Week vs week", "Spikes", "Mentions"} <= set(workbook.sheetnames)
    about = {row[0].value: row[1].value for row in workbook["About"].iter_rows(min_row=2)}
    assert "fictional" in about["Demo notice"] and "z-score" in about["Spike rule"]
    html = pulse_html(report)
    assert "Fictional demo" in html and "not a verdict" in html and "threshold" in html
    assert html.count("<html") == 1 and "<script" not in html


def test_spreadsheet_formula_injection_is_neutralized():
    assert safe_cell("=HYPERLINK(\"http://x\")") == "'=HYPERLINK(\"http://x\")"
    assert safe_cell("- Jeg har nesten gitt opp") == "'- Jeg har nesten gitt opp"
    assert safe_cell("+47 99") == "'+47 99" and safe_cell("@SUM(A1)") == "'@SUM(A1)"
    assert safe_cell("Vanlig tittel\x07") == "Vanlig tittel" and safe_cell(3) == 3


def test_html_escapes_feed_text():
    ws = demo_workspace()
    articles = ws.articles.copy()
    last = articles.index[-1]
    articles.loc[last, "title"] = "<script>alert(1)</script> Fjellbrus"
    # newest item of the final demo day, so it is listed among the latest mentions
    articles.loc[last, "published"] = pd.Timestamp(DEMO_END).tz_localize("Europe/Oslo") + timedelta(hours=23, minutes=59)
    poisoned = type(ws)(ws.label, ws.is_demo, ws.brands, articles, mention_table(articles, ws.brands), ws.fetch_log, ws.notice)
    html = pulse_html(build_pulse(poisoned, DEMO_END))
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html
